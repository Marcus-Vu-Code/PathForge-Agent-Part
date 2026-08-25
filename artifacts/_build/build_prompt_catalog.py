from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(r"F:\Github\PathForge Agent Part\artifacts\pathforge_prompt_catalog.docx")


PROMPTS = [
    {
        "name": "Backend extraction system prompt",
        "source": "backend/app/providers/openai_provider.py:18",
        "used_by": "OpenAI, Gemini, and Qwen background extraction providers",
        "purpose": "Turn uploaded background artifacts into a typed BackgroundArtifactExtraction / CareerProfile.",
        "prompt": (
            "Extract a versioned CareerProfile from the submitted background.\n"
            "Use only facts present in the artifact. Preserve uncertainty in warnings. Return JSON only."
        ),
        "inputs": "UploadedArtifact JSON: artifact type, filename, content type, and extracted text.",
        "output": "BackgroundArtifactExtraction schema with entities, evidence spans, confidence, and warnings.",
    },
    {
        "name": "Backend reasoning system prompt",
        "source": "backend/app/providers/openai_provider.py:22",
        "used_by": "OpenAI, Gemini, and Qwen career reasoners",
        "purpose": "Synthesize profile, goal, deterministic tools, and evidence into a typed CareerPlan.",
        "prompt": (
            "You are PathForge AI, a career-navigation reasoner.\n"
            "Synthesize the supplied profile, goal, deterministic tool outputs, and evidence into a typed CareerPlan.\n"
            "Do not invent user skills or market facts. Use caveats for uncertainty. Return JSON only."
        ),
        "inputs": "CareerReasoningRequest JSON: profile, goal, evidence, gaps, project recommendations, and next actions.",
        "output": "CareerPlan schema with paths, strengths, gaps, actions, projects, evidence, caveats, and confidence.",
    },
    {
        "name": "Qwen extraction schema wrapper",
        "source": "backend/app/providers/qwen_provider.py:66",
        "used_by": "Local Qwen background extractor",
        "purpose": "Add an explicit JSON Schema requirement to the shared extraction prompt.",
        "prompt": (
            "{EXTRACTION_PROMPT}\n"
            "Return one JSON object that validates against this JSON schema:\n"
            "{BackgroundArtifactExtraction JSON schema}"
        ),
        "inputs": "Potentially truncated UploadedArtifact JSON, limited by QWEN_MAX_INPUT_CHARS.",
        "output": "One JSON object that validates as BackgroundArtifactExtraction.",
    },
    {
        "name": "Qwen reasoning schema wrapper",
        "source": "backend/app/providers/qwen_provider.py:92",
        "used_by": "Local Qwen career reasoner",
        "purpose": "Add an explicit JSON Schema requirement to the shared reasoning prompt.",
        "prompt": (
            "{REASONING_PROMPT}\n"
            "Return one JSON object only. It must validate against this JSON schema:\n"
            "{CareerPlan JSON schema}"
        ),
        "inputs": "CareerReasoningRequest JSON.",
        "output": "One JSON object that validates as CareerPlan.",
    },
    {
        "name": "Hosted extraction system prompt",
        "source": "frontend/worker/index.js:619",
        "used_by": "Vercel hosted API route for extraction",
        "purpose": "Lightweight hosted extraction prompt for same-origin API functions.",
        "prompt": (
            "Extract a career profile JSON object with keys: education, experience, skills, projects, "
            "certifications, interests, constraints, source_artifacts, extraction_confidence, warnings."
        ),
        "inputs": "Uploaded or pasted background text, including text extracted from DOCX/PDF where possible.",
        "output": "Profile-like JSON normalized by the hosted worker.",
    },
    {
        "name": "Hosted reasoning system prompt",
        "source": "frontend/worker/index.js:623",
        "used_by": "Vercel hosted API route for Run Agent",
        "purpose": "Lightweight hosted reasoning prompt for Gemini/OpenAI-compatible chat completions.",
        "prompt": (
            "Create a career plan JSON object with keys: recommended_paths, strengths, gaps, next_actions, "
            "project_recommendations, evidence, caveats, overall_confidence. Keep every recommendation practical "
            "and evidence-grounded."
        ),
        "inputs": "Profile, goal, evidence, gaps, projects, and next_actions serialized as JSON.",
        "output": "Plan-like JSON normalized by the hosted worker before rendering.",
    },
    {
        "name": "Hosted JSON-only instruction",
        "source": "frontend/worker/index.js:279",
        "used_by": "Hosted Gemini/OpenAI-compatible chat completions",
        "purpose": "Force model responses toward JSON-only output for parser reliability.",
        "prompt": "{systemPrompt}\nReturn only valid JSON.",
        "inputs": "App-selected system prompt plus user/background payload.",
        "output": "JSON text parsed by the hosted worker.",
    },
    {
        "name": "Background prompt template",
        "source": "backend/app/services/document_intake.py:101 and frontend/worker/index.js:474",
        "used_by": "Build Background Prompt button",
        "purpose": "Turn extracted profile fields into a reusable prompt for career-fit evaluation.",
        "prompt": (
            "Use the background below to evaluate career fit, identify role-specific evidence, and recommend practical next steps.\n\n"
            "Candidate Background\n"
            "{Education}\n{Experience}\n{Projects}\n{Skills}\n{Certifications}\n{Interests}\n{Constraints}\n\n"
            "Source Documents\n"
            "- {filename} ({artifact_type})\n\n"
            "Review Notes\n"
            "- {warnings}"
        ),
        "inputs": "BackgroundArtifactExtraction plus document metadata.",
        "output": "Plain-text candidate background prompt shown in the app.",
    },
    {
        "name": "Guided question: current situation",
        "source": "frontend/src/main.tsx:30",
        "used_by": "Answer Background Questions flow",
        "purpose": "Collect current context from users without a resume.",
        "prompt": "What are you doing now?",
        "inputs": "Free text.",
        "output": "Background section appended to the profile extraction text.",
    },
    {
        "name": "Guided question: education and training",
        "source": "frontend/src/main.tsx:38",
        "used_by": "Answer Background Questions flow",
        "purpose": "Collect education, training, certifications, and self-study.",
        "prompt": "What school, training, courses, certifications, or self-study have you completed?",
        "inputs": "Free text.",
        "output": "Education section appended to the background text.",
    },
    {
        "name": "Guided question: work and responsibilities",
        "source": "frontend/src/main.tsx:46",
        "used_by": "Answer Background Questions flow",
        "purpose": "Collect work, volunteer, internship, and responsibility evidence.",
        "prompt": "What jobs, volunteer work, internships, or responsibilities have you had?",
        "inputs": "Free text.",
        "output": "Experience section appended to the background text.",
    },
    {
        "name": "Guided question: skills and tools",
        "source": "frontend/src/main.tsx:54",
        "used_by": "Answer Background Questions flow",
        "purpose": "Collect tool, software, language, and strength evidence.",
        "prompt": "What skills, tools, software, languages, or strengths do you have?",
        "inputs": "Free text.",
        "output": "Skills section appended to the background text.",
    },
    {
        "name": "Guided question: projects and accomplishments",
        "source": "frontend/src/main.tsx:62",
        "used_by": "Answer Background Questions flow",
        "purpose": "Collect project, accomplishment, process improvement, and portfolio evidence.",
        "prompt": "What have you built, improved, organized, solved, or helped complete?",
        "inputs": "Free text.",
        "output": "Projects section appended to the background text.",
    },
    {
        "name": "Guided question: interests",
        "source": "frontend/src/main.tsx:70",
        "used_by": "Answer Background Questions flow",
        "purpose": "Capture field, industry, and problem-interest signals.",
        "prompt": "What kind of work, industries, or problems are you interested in?",
        "inputs": "Free text.",
        "output": "Interests section appended to the background text.",
    },
    {
        "name": "Guided question: constraints and preferences",
        "source": "frontend/src/main.tsx:78",
        "used_by": "Answer Background Questions flow",
        "purpose": "Capture planning constraints such as time, location, remote/hybrid, budget, schedule, and dealbreakers.",
        "prompt": "What should the plan account for?",
        "inputs": "Free text.",
        "output": "Constraints section appended to the background text.",
    },
    {
        "name": "Guided question: anything else",
        "source": "frontend/src/main.tsx:86",
        "used_by": "Answer Background Questions flow",
        "purpose": "Let users add context that does not fit the structured questions.",
        "prompt": "What else should the agent know about you?",
        "inputs": "Free text.",
        "output": "Additional context section appended to the background text.",
    },
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_font(run, name="Calibri", size=11, color="000000", bold=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold


def add_prompt_block(doc: Document, index: int, item: dict[str, str]) -> None:
    heading = doc.add_heading(f"{index}. {item['name']}", level=1)
    heading.runs[0].font.color.rgb = RGBColor.from_string("2E74B5")

    meta = doc.add_table(rows=4, cols=2)
    meta.style = "Table Grid"
    rows = [
        ("Source", item["source"]),
        ("Used by", item["used_by"]),
        ("Purpose", item["purpose"]),
        ("Output", item["output"]),
    ]
    for row, (label, value) in zip(meta.rows, rows):
        label_cell, value_cell = row.cells
        label_cell.text = label
        value_cell.text = value
        set_cell_width(label_cell, 1800)
        set_cell_width(value_cell, 7560)
        set_cell_shading(label_cell, "F2F4F7")
        for paragraph in label_cell.paragraphs:
            for run in paragraph.runs:
                set_font(run, bold=True)
        for paragraph in value_cell.paragraphs:
            for run in paragraph.runs:
                set_font(run)

    p = doc.add_paragraph()
    label = p.add_run("Prompt text")
    set_font(label, bold=True, color="1F4D78")

    prompt_table = doc.add_table(rows=1, cols=1)
    prompt_table.style = "Table Grid"
    cell = prompt_table.cell(0, 0)
    set_cell_width(cell, 9360)
    set_cell_shading(cell, "F8FAFC")
    cell.text = item["prompt"]
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            set_font(run, name="Consolas", size=9)

    p = doc.add_paragraph()
    r = p.add_run(f"Inputs: {item['inputs']}")
    set_font(r, size=10, color="565D66")


def main() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(11)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run("PathForge Prompt Catalog")
    set_font(run, size=24, color="0B2545", bold=True)

    subtitle = doc.add_paragraph()
    r = subtitle.add_run("Implementation prompts, prompt templates, and user-facing guided questions gathered from the current PathForge codebase.")
    set_font(r, size=11, color="565D66")

    doc.add_heading("How to use this document", level=1)
    for text in [
        "Use this as a reference when explaining, debugging, or revising PathForge behavior.",
        "Provider prompts are schema-bound; changing the schemas or prompt wording can affect extraction and plan rendering.",
        "Do not paste real resumes, API keys, or sensitive user documents into this catalog.",
    ]:
        p = doc.add_paragraph(style="List Bullet")
        set_font(p.add_run(text), size=11)

    doc.add_heading("Prompt inventory", level=1)
    summary = doc.add_table(rows=1, cols=4)
    summary.style = "Table Grid"
    headers = ["Prompt", "Source", "Used by", "Purpose"]
    for idx, header in enumerate(headers):
        cell = summary.cell(0, idx)
        cell.text = header
        set_cell_shading(cell, "E8EEF5")
        for run in cell.paragraphs[0].runs:
            set_font(run, bold=True)
    widths = [2500, 2300, 2200, 2360]
    for item in PROMPTS:
        row = summary.add_row().cells
        values = [item["name"], item["source"], item["used_by"], item["purpose"]]
        for idx, value in enumerate(values):
            row[idx].text = value
            set_cell_width(row[idx], widths[idx])
            for paragraph in row[idx].paragraphs:
                for run in paragraph.runs:
                    set_font(run, size=9)

    doc.add_section(WD_SECTION.NEW_PAGE)
    for index, item in enumerate(PROMPTS, start=1):
        add_prompt_block(doc, index, item)

    doc.save(OUT)


if __name__ == "__main__":
    main()
